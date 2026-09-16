"""`tools/marcar_plantilla.py` · los campos globales de portada y cabecera.

**Este fichero no existía**, y es el motivo por el que lo que se prueba aquí
llegó a estar roto en tres de las cuatro plantillas del cliente sin que nadie se
enterara. Se descubrió pasando las otras tres por la herramienta y **contando
los marcadores** del resultado, no leyendo el código: el Modelo A castellano
—la única con la que se venía trabajando— salía perfecta, y las otras tres
perdían la portada, la cabecera o las dos.

Las plantillas de prueba se **construyen aquí**, con python-pptx: ninguna del
cliente entra en el repositorio. Lo que se comprueba no es que sepa abrir un
PPTX real —eso lo hace la biblioteca— sino que reconoce **las variantes de
redacción que las cuatro plantillas usan de verdad**, que es donde fallaba.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path
from types import ModuleType

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[4]
HERRAMIENTA = RAIZ / "tools" / "marcar_plantilla.py"


def _cargar() -> ModuleType:
    """El programa vive en `tools/`, que no es un paquete instalable."""
    spec = importlib.util.spec_from_file_location("marcar_plantilla", HERRAMIENTA)
    assert spec is not None and spec.loader is not None
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = modulo
    spec.loader.exec_module(modulo)
    return modulo


herramienta = _cargar()

VACIA = 6  # la disposición «en blanco» de la plantilla por defecto de python-pptx


def _caja(slide, textos: list[str], *, arriba: float = 0.5):
    """Un cuadro de texto con un párrafo por línea."""
    caja = slide.shapes.add_textbox(Inches(0.4), Inches(arriba), Inches(8.5), Inches(1.2))
    marco = caja.text_frame
    for i, texto in enumerate(textos):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        run = parrafo.add_run()
        run.text = texto
        run.font.size = Pt(18)
    return caja


def _portada(titulo: str, nombre: str, fecha: str) -> Presentation:
    prs = Presentation()
    _caja(prs.slides.add_slide(prs.slide_layouts[VACIA]), [titulo, nombre, fecha])
    return prs


def _textos(prs: Presentation) -> list[str]:
    return [
        p.text.strip()
        for s in prs.slides
        for f in s.shapes
        if getattr(f, "text_frame", None) is not None
        for p in f.text_frame.paragraphs
        if p.text.strip()
    ]


# ─────────────────────────────────────────────────────────────────────────────
#  La portada
# ─────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("titulo", "nombre", "fecha"),
    [
        # Modelo A castellano: relleno de equis y una fecha DE VERDAD, de otro
        # encargo. Es la única de las cuatro que está así.
        ("Due Dilligence Técnica", "XXXXX", "Febrero 2026"),
        # Modelo B castellano: el nombre del campo, no un relleno.
        ("Due Dilligence Técnica", "PROYECTO", "FECHA"),
        # Las dos inglesas. Y con la errata: «Dilligence», con dos eles.
        ("Technical Due Dilligence", "PROJECT", "Date"),
    ],
)
def test_la_portada_se_marca_escriba_como_escriba_sus_huecos(
    titulo: str, nombre: str, fecha: str
) -> None:
    """`[REQ]` Las cuatro plantillas rellenan la portada de tres maneras y solo
    una estaba contemplada. Las otras dos salían con la palabra «PROJECT» por
    título del informe y sin fecha, que es de lo primero que se ve."""
    prs = _portada(titulo, nombre, fecha)
    herramienta.marcar(prs)
    textos = _textos(prs)
    assert "{{project.name}}" in textos
    assert "{{report.month}}" in textos


def test_la_errata_del_titulo_no_deja_la_portada_sin_marcar() -> None:
    """Las cuatro escriben «Dilligence» con dos eles, también las inglesas. La
    errata estaba contemplada **solo en castellano**, así que las dos inglesas
    no se reconocían como portada y se quedaban enteras sin marcar."""
    prs = _portada("Technical Due Dilligence", "PROJECT", "Date")
    herramienta.marcar(prs)
    assert "{{project.name}}" in _textos(prs)

    # Y la grafía correcta sigue valiendo: no se ha cambiado una por otra.
    prs = _portada("Technical Due Diligence", "PROJECT", "Date")
    herramienta.marcar(prs)
    assert "{{project.name}}" in _textos(prs)


def test_una_primera_diapositiva_que_no_es_portada_no_se_toca() -> None:
    """La comprobación del título es lo que impide que cualquier documento cuya
    primera página lleve la palabra «PROJECT» acabe con el nombre del cliente
    metido donde no va."""
    prs = Presentation()
    _caja(prs.slides.add_slide(prs.slide_layouts[VACIA]), ["INDICE", "PROJECT", "Date"])
    herramienta.marcar(prs)
    assert _textos(prs) == ["INDICE", "PROJECT", "Date"]


def test_el_resalte_amarillo_del_hueco_no_llega_al_informe() -> None:
    """`[REQ]` Tres de las cuatro escriben su «PROJECT» **resaltado en
    amarillo**: es la marca de «esto hay que rellenarlo», igual que las equis de
    la castellana, no una decisión de diseño. Conservarlo sacaba la portada con
    el nombre del cliente subrayado en fosforito, y así se vio al renderizar.

    El formato **por lo demás se conserva** —cuerpo, color, tipografía—: es la
    regla de la casa con la plantilla de un cliente, y esta es la excepción
    justificada, no una licencia para reformatear.
    """
    from pptx.dml.color import RGBColor

    prs = _portada("Technical Due Dilligence", "PROJECT", "Date")
    parrafo = prs.slides[0].shapes[0].text_frame.paragraphs[1]
    run = parrafo.runs[0]
    rpr = run._r.get_or_add_rPr()
    resalte = rpr.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}highlight", {}
    )
    color = rpr.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr", {"val": "FFFF00"}
    )
    resalte.append(color)
    rpr.append(resalte)
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)

    herramienta.marcar(prs)

    marcado = prs.slides[0].shapes[0].text_frame.paragraphs[1].runs[0]
    assert marcado.text == "{{project.name}}"
    assert "highlight" not in marcado._r.xml
    assert marcado.font.size == Pt(28)
    assert marcado.font.color.rgb == RGBColor(0x11, 0x22, 0x33)


# ─────────────────────────────────────────────────────────────────────────────
#  La cabecera de todas las páginas
# ─────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("rotulo", ["NOMBRE DEL PROYECTO", "NOMBRE PROYECTO", "PROJECT NAME"])
def test_el_rotulo_de_cabecera_se_reconoce_en_sus_tres_grafias(rotulo: str) -> None:
    """`[REQ]` «NOMBRE PROYECTO», sin el «DEL», es el del Modelo B castellano y
    no estaba. Sin él esa plantilla se quedaba sin un solo `{{project.name}}` en
    sus patrones: once páginas con el rótulo literal en la cabecera, que es
    exactamente como salió el primer informe que se generó con ella."""
    prs = Presentation()
    _caja(prs.slides.add_slide(prs.slide_layouts[VACIA]), ["ANÁLISIS TÉCNICO", rotulo])
    herramienta.marcar(prs)
    assert "{{project.name}}" in _textos(prs)


def test_el_titulo_de_la_seccion_no_se_toca() -> None:
    """Va en el párrafo de arriba del mismo cuadro y es **contenido del
    informe**, no un dato del proyecto: marcarlo dejaría la sección sin nombre."""
    prs = Presentation()
    _caja(
        prs.slides.add_slide(prs.slide_layouts[VACIA]),
        ["ESTIMACIÓN ECONÓMICA - CAPEX", "NOMBRE DEL PROYECTO"],
    )
    herramienta.marcar(prs)
    assert "ESTIMACIÓN ECONÓMICA - CAPEX" in _textos(prs)


# ─────────────────────────────────────────────────────────────────────────────
#  Los títulos que cada plantilla escribe a su manera
# ─────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("titulo", ["EMPLAZAMIENTO", "LOCATION", "PLOT LOCATION"])
def test_la_direccion_se_reconoce_en_las_cuatro(titulo: str) -> None:
    """Las dos inglesas titulan este cuadro «PLOT LOCATION», y `LOCATION` a
    secas no aparece en ninguna de las cuatro: estaba puesto de memoria."""
    prs = Presentation()
    _caja(prs.slides.add_slide(prs.slide_layouts[VACIA]), [titulo, "XXXXXXXXXXXX"])
    herramienta.marcar(prs)
    assert "{{asset.address}}" in _textos(prs)


@pytest.mark.parametrize("titulo", ["LICENSE ANALYSIS", "LICENSES ANALYSIS"])
def test_el_analisis_de_licencias_en_singular_y_en_plural(titulo: str) -> None:
    """Cuatro plantillas, cuatro grafías: el Modelo A inglés lo escribe en
    singular y el B inglés en plural. Una letra de diferencia dejaba la sección
    sin su marcador."""
    assert titulo in herramienta.CAMPOS_POR_PATRON
    assert herramienta.CAMPOS_POR_PATRON[titulo] == "{{docs.licencias}}"


def test_un_titulo_que_no_se_conoce_no_se_inventa() -> None:
    """Mejor un hueco con su «XXXX» a la vista —que quien redacta corrige— que
    un marcador puesto a ojo en el sitio equivocado."""
    prs = Presentation()
    _caja(prs.slides.add_slide(prs.slide_layouts[VACIA]), ["SECCIÓN INVENTADA", "XXXXXXXXXXXX"])
    herramienta.marcar(prs)
    assert "XXXXXXXXXXXX" in _textos(prs)
    assert not [t for t in _textos(prs) if t.startswith("{{")]
