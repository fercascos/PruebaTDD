"""El rótulo que cae con su marcador `[REQ]` `@rotulo`.

El cliente lo pidió al ver el informe generado: *«quita el rótulo si no hay
valoración»*. Su plantilla escribe «Valoración» en un párrafo aparte, encima del
hueco, y ese párrafo **no es un marcador**: cuando no había valoración, el hueco
se vaciaba —como debe— y el rótulo se quedaba solo, encabezando media página en
blanco.

La plantilla de prueba se construye aquí: ninguna del cliente entra en el
repositorio.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from tdd.reporting import composicion
from tdd.reporting.clone import sustituir_marcadores

VACIA = 6  # la disposición «en blanco» de la plantilla por defecto de python-pptx


def _diapositiva(prs: Presentation, parrafos: list[str], notas: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[VACIA])
    marco = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(4)).text_frame
    for i, texto in enumerate(parrafos):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.add_run().text = texto
    slide.notes_slide.notes_text_frame.text = notas


def _pagina(valores: dict[str, str]) -> list[str]:
    """Genera la página de CUBIERTA con esos valores y devuelve su texto."""
    prs = Presentation()
    _diapositiva(
        prs,
        [
            "CUBIERTA",
            "{{descriptivo:HC.H02}}",
            "{{rotulo:valoracion:HC.H02}}",
            "{{valoracion:HC.H02}}",
        ],
        "@rotulo: valoracion:HC.H02 = Valoración",
    )
    composicion.resolver_rotulos(valores, composicion.rotulos_de(prs))
    for slide in prs.slides:
        sustituir_marcadores(slide, valores)
    return [
        p.text.strip()
        for s in prs.slides
        for f in s.shapes
        if getattr(f, "text_frame", None) is not None
        for p in f.text_frame.paragraphs
    ]


def test_con_valoracion_el_rotulo_se_queda() -> None:
    texto = _pagina(
        {
            "descriptivo:HC.H02": "Cubierta deck con lámina de PVC.",
            "valoracion:HC.H02": "Lámina al final de su vida útil.",
        }
    )
    assert texto == [
        "CUBIERTA",
        "Cubierta deck con lámina de PVC.",
        "Valoración",
        "Lámina al final de su vida útil.",
    ]


def test_sin_valoracion_el_rotulo_se_va_con_ella() -> None:
    """Lo que el cliente pidió. Y lo que queda no es «Valoración» y un hueco:
    es el descriptivo y nada más, que es una página correcta."""
    texto = _pagina({"descriptivo:HC.H02": "Cubierta deck con lámina de PVC."})
    assert texto == ["CUBIERTA", "Cubierta deck con lámina de PVC.", "", ""]


def test_una_valoracion_en_blanco_cuenta_como_no_haberla() -> None:
    """Un texto de solo espacios llega de un campo que alguien tocó y dejó
    igual. Si contara como valoración, el rótulo volvería a quedarse solo."""
    texto = _pagina({"descriptivo:HC.H02": "Cubierta deck.", "valoracion:HC.H02": "   "})
    assert "Valoración" not in texto


def test_el_rotulo_no_imprime_nunca_su_marcador() -> None:
    """`[REQ]` §17.7 · Ni con valoración ni sin ella puede salir un `{{...}}`
    literal en la pantalla del cliente."""
    for valores in ({"valoracion:HC.H02": "Algo."}, {}):
        assert not [t for t in _pagina(dict(valores)) if "{{" in t]


def test_el_texto_del_rotulo_sale_de_la_plantilla_y_no_de_una_constante() -> None:
    """Es palabra del cliente y cambia con el idioma: sus plantillas castellanas
    ponen «Valoración» y las inglesas **«Valuation»**. Escribirla en el código
    sería traducirle el informe sin permiso."""
    prs = Presentation()
    _diapositiva(
        prs,
        ["ROOF", "{{rotulo:valoracion:HC.H02}}", "{{valoracion:HC.H02}}"],
        "@rotulo: valoracion:HC.H02 = Valuation",
    )
    valores = {"valoracion:HC.H02": "End of service life."}
    composicion.resolver_rotulos(valores, composicion.rotulos_de(prs))
    assert valores["rotulo:valoracion:HC.H02"] == "Valuation"


def test_las_notas_de_una_diapositiva_no_afectan_a_otra() -> None:
    """Los rótulos se leen de toda la presentación —una plantilla tiene catorce,
    uno por sección— y cada uno se ata a **su** marcador, no a la diapositiva en
    la que estaba escrito."""
    prs = Presentation()
    _diapositiva(prs, ["A"], "@rotulo: valoracion:HC.H02 = Valoración")
    _diapositiva(prs, ["B"], "@rotulo: valoracion:HC.H03 = Valoración")
    valores = {"valoracion:HC.H03": "Sellado de juntas endurecido."}
    composicion.resolver_rotulos(valores, composicion.rotulos_de(prs))
    assert valores["rotulo:valoracion:HC.H02"] == ""
    assert valores["rotulo:valoracion:HC.H03"] == "Valoración"
