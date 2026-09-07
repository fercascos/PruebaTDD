"""`tools/retipografiar_plantilla.py` · el cambio de tipografía de P-39.

La plantilla de prueba se **construye aquí**, con python-pptx: ninguna del
cliente entra en el repositorio. Lo que se comprueba no es que sepa abrir un
PPTX real —eso lo hace la biblioteca— sino las tres decisiones del programa:
que empareja por peso, que no toca el original y que señala lo que no sabe
traducir en vez de adivinarlo.
"""

from __future__ import annotations

import importlib.util
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from types import ModuleType

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[4]
HERRAMIENTA = RAIZ / "tools" / "retipografiar_plantilla.py"


def _cargar() -> ModuleType:
    """El programa vive en `tools/`, que no es un paquete instalable."""
    spec = importlib.util.spec_from_file_location("retipografiar_plantilla", HERRAMIENTA)
    assert spec is not None and spec.loader is not None
    modulo = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = modulo
    spec.loader.exec_module(modulo)
    return modulo


herramienta = _cargar()


def _plantilla(tmp_path: Path, fuentes: dict[str, str]) -> Path:
    """Un PPTX con un cuadro de texto por familia pedida."""
    pres = Presentation()
    diapositiva = pres.slides.add_slide(pres.slide_layouts[6])
    for i, (familia, texto) in enumerate(fuentes.items()):
        caja = diapositiva.shapes.add_textbox(Inches(0.5), Inches(0.5 + i), Inches(8), Inches(0.8))
        run = caja.text_frame.paragraphs[0].add_run()
        run.text = texto
        run.font.name = familia
        run.font.size = Pt(12)
    destino = tmp_path / "plantilla.pptx"
    pres.save(destino)
    return destino


#: `<a:majorFont><a:latin typeface="…">` para titulares, `minorFont` para el
#: cuerpo. python-pptx no expone el tema, así que se edita el XML.
_TEMA = re.compile(rb'(<a:(?:majorFont|minorFont)>\s*<a:latin typeface=")[^"]*(")')


def _poner_tema(bruto: bytes, titulares: bytes, cuerpo: bytes) -> bytes:
    # `majorFont` va antes que `minorFont` en el XML del tema, así que el orden
    # de consumo es el de aparición.
    orden = iter([titulares, cuerpo])
    return _TEMA.sub(lambda m: m.group(1) + next(orden) + m.group(2), bruto)


def _con_tema(plantilla: Path, titulares: bytes, cuerpo: bytes) -> Path:
    """La misma plantilla, con el TEMA declarando esas dos familias."""
    with zipfile.ZipFile(plantilla) as z:
        nombres = z.namelist()
        partes = {n: z.read(n) for n in nombres}
    for nombre in nombres:
        if nombre.startswith("ppt/theme/"):
            partes[nombre] = _poner_tema(partes[nombre], titulares, cuerpo)
    with zipfile.ZipFile(plantilla, "w", zipfile.ZIP_DEFLATED) as salida:
        for n in nombres:
            salida.writestr(n, partes[n])
    return plantilla


#: Lo mismo, pero capturando el nombre para poder leerlo.
_TEMA_LEER = re.compile(rb'<a:(?:majorFont|minorFont)>\s*<a:latin typeface="([^"]*)"')


def _tema_de(fichero: Path) -> list[str]:
    """[titulares, cuerpo] tal y como los declara el tema."""
    with zipfile.ZipFile(fichero) as z:
        tema = z.read("ppt/theme/theme1.xml")
    return [m.decode() for m in _TEMA_LEER.findall(tema)]


def test_el_tema_tambien_se_convierte(tmp_path: Path) -> None:
    """**La parte que más pesa en una plantilla de verdad.**

    Una corporativa fija su tipografía en el tema —`majorFont` para titulares y
    `minorFont` para el cuerpo—, y todo lo que no lleva formato explícito la
    hereda de ahí. Convertir solo los `run` dejaría el grueso del informe
    apuntando a una fuente que en el PowerPoint del lector no existe, que es
    exactamente el fallo silencioso que esta herramienta viene a evitar.
    """
    origen = _con_tema(
        _plantilla(tmp_path, {"Gotham Light": "cuerpo"}), b"Gotham Ultra", b"Gotham Light"
    )
    assert _tema_de(origen) == ["Gotham Ultra", "Gotham Light"]

    salida = tmp_path / "nueva.pptx"
    herramienta.retipografiar(origen, salida, mapa=herramienta.EQUIVALENCIAS)
    assert _tema_de(salida) == ["Montserrat Black", "Montserrat Light"]


def test_declara_lo_que_hay_dentro(tmp_path: Path) -> None:
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo", "Gotham Ultra": "TITULAR"})
    cuenta = herramienta.declaradas(origen)
    assert cuenta["Gotham Light"] >= 1
    assert cuenta["Gotham Ultra"] >= 1


def test_empareja_por_peso_y_no_por_prefijo(tmp_path: Path) -> None:
    """El fallo que la regla evita: con «empieza por Gotham», `Gotham Ultra`
    habría casado con la regla de `Gotham` a secas y **todos los titulares del
    informe** habrían acabado en el peso del cuerpo."""
    origen = _plantilla(
        tmp_path,
        {"Gotham Light": "cuerpo", "Gotham Ultra": "TITULAR", "Gotham Medium": "cabecera"},
    )
    salida = tmp_path / "nueva.pptx"
    cambios, sin_traducir = herramienta.retipografiar(
        origen, salida, mapa=herramienta.EQUIVALENCIAS
    )

    resultado = herramienta.declaradas(salida)
    assert resultado["Montserrat Light"] >= 1
    assert resultado["Montserrat Black"] >= 1  # Ultra es el peso de titular
    assert resultado["Montserrat Medium"] >= 1
    assert not any(f.startswith("Gotham") for f in resultado)
    assert not sin_traducir
    assert isinstance(cambios, Counter)


def test_el_original_no_se_toca(tmp_path: Path) -> None:
    """`[REQ]` Regla de la casa: los ficheros que trae el cliente no se
    sobrescriben nunca."""
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo"})
    antes = origen.read_bytes()
    herramienta.retipografiar(origen, tmp_path / "nueva.pptx", mapa=herramienta.EQUIVALENCIAS)
    assert origen.read_bytes() == antes


def test_escribir_sobre_el_original_se_rechaza_incluso_forzando(tmp_path: Path) -> None:
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo"})
    with pytest.raises(SystemExit, match="original"):
        herramienta.retipografiar(origen, origen, mapa=herramienta.EQUIVALENCIAS, forzar=True)


def test_no_pisa_un_fichero_existente_sin_permiso(tmp_path: Path) -> None:
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo"})
    ocupado = tmp_path / "ocupado.pptx"
    ocupado.write_bytes(b"contenido previo")
    with pytest.raises(SystemExit, match="forzar"):
        herramienta.retipografiar(origen, ocupado, mapa=herramienta.EQUIVALENCIAS)
    assert ocupado.read_bytes() == b"contenido previo"


def test_una_familia_desconocida_se_señala_y_se_deja(tmp_path: Path) -> None:
    """Mejor una familia sin traducir y señalada que una traducida a ojo: el
    peso equivocado en un titular no lo ve nadie hasta que está impreso."""
    origen = _plantilla(tmp_path, {"Gotham Rounded Book": "raro", "Gotham Light": "cuerpo"})
    salida = tmp_path / "nueva.pptx"
    _, sin_traducir = herramienta.retipografiar(origen, salida, mapa=herramienta.EQUIVALENCIAS)
    assert sin_traducir == ["Gotham Rounded Book"]
    assert herramienta.declaradas(salida)["Gotham Rounded Book"] >= 1


def test_lo_que_no_es_tipografia_se_copia_byte_a_byte(tmp_path: Path) -> None:
    """La plantilla conserva sus imágenes, sus gráficos y su diseño: solo
    cambian los atributos `typeface` de las partes XML."""
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo"})
    salida = tmp_path / "nueva.pptx"
    herramienta.retipografiar(origen, salida, mapa=herramienta.EQUIVALENCIAS)

    with zipfile.ZipFile(origen) as a, zipfile.ZipFile(salida) as b:
        assert a.namelist() == b.namelist()
        distintas = [n for n in a.namelist() if a.read(n) != b.read(n)]
    # Solo la diapositiva que llevaba la fuente.
    assert all(n.endswith(".xml") for n in distintas)
    assert len(distintas) == 1


def test_century_gothic_solo_si_se_pide(tmp_path: Path) -> None:
    """Venía de las tablas pegadas desde Excel (P-38). No se traduce por
    omisión: en una plantilla puede ser una decisión de diseño y no un resto."""
    origen = _plantilla(tmp_path, {"Century Gothic": "tabla vieja"})
    sin, con = tmp_path / "sin.pptx", tmp_path / "con.pptx"

    herramienta.retipografiar(origen, sin, mapa=herramienta.EQUIVALENCIAS)
    assert herramienta.declaradas(sin)["Century Gothic"] >= 1

    mapa = {**herramienta.EQUIVALENCIAS, **herramienta.EQUIVALENCIAS_EXTRA}
    herramienta.retipografiar(origen, con, mapa=mapa)
    assert "Century Gothic" not in herramienta.declaradas(con)


def test_es_idempotente(tmp_path: Path) -> None:
    origen = _plantilla(tmp_path, {"Gotham Light": "cuerpo"})
    una, dos = tmp_path / "1.pptx", tmp_path / "2.pptx"
    herramienta.retipografiar(origen, una, mapa=herramienta.EQUIVALENCIAS)
    cambios, _ = herramienta.retipografiar(una, dos, mapa=herramienta.EQUIVALENCIAS)
    assert sum(cambios.values()) == 0
    assert herramienta.declaradas(dos) == herramienta.declaradas(una)


def test_las_equivalencias_apuntan_a_familias_que_el_informe_declara() -> None:
    """Traducir a una familia que el informe no exige dejaría la plantilla
    apuntando a algo que la imagen no instala."""
    from tdd.reporting.fonts import FAMILIAS_REQUERIDAS

    # Thin y ExtraLight existen en Montserrat pero el informe no los exige: no
    # se usan en las plantillas y no tiene sentido bloquear el arranque por
    # ellos. El resto sí tiene que estar entre las requeridas.
    sueltas = {"Montserrat Thin", "Montserrat ExtraLight"}
    destinos = set(herramienta.EQUIVALENCIAS.values()) - sueltas
    assert destinos <= set(FAMILIAS_REQUERIDAS)
