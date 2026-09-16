"""Diapositivas que se repiten `[REQ]` §17.2, regla 5 · `@repeat`.

El cliente lo pidió con estas palabras al hablar del Full Report: *«que vayas
rellenando con los textos que aparecen en la aplicación las distintas slides del
informe»*. Un proyecto de cartera tiene seis activos y cuarenta hallazgos, y una
plantilla tiene **una** diapositiva de activo y **una** de hallazgo: alguien las
duplicaba a mano en PowerPoint y volvía a teclear los datos en cada copia.

## Cómo se declara

En las **notas del orador** de la diapositiva modelo, una línea:

    @repeat: asset
    @repeat: finding
    @max: 20

Van en las notas y no en el cuerpo porque **no se imprimen**: la diapositiva se
ve exactamente igual en la plantilla y en el informe, y la instrucción no se
cuela en un entregable si alguien se olvida de borrarla.

## Qué hace

La diapositiva modelo se **clona una vez por elemento**, cada copia recibe los
marcadores de *su* elemento —`{{asset.name}}` es el nombre de ese activo, no el
del primero— y el modelo se retira al final. Las copias van **en el sitio del
modelo**, no al final de la presentación: un informe cuyas diapositivas de
activo aparecen detrás de las conclusiones no es el informe que se diseñó.

`[LIM]` **La colección vacía deja la diapositiva fuera.** Un proyecto sin
hallazgos no produce una diapositiva de hallazgo en blanco con los marcadores a
la vista, que es lo que delata un informe hecho a máquina. Si lo que se quería
era justamente ver el hueco, se quita el `@repeat` y la diapositiva se queda
fija.
"""

from __future__ import annotations

import re
from typing import Any

from pptx.presentation import Presentation as Presentacion
from pptx.slide import Slide

from tdd.reporting import marcadores as mk
from tdd.reporting.clone import clonar_diapositiva

#: `@repeat: asset`, con espacios a gusto de quien lo escriba.
REPETIR = re.compile(r"@repeat\s*:\s*([a-zA-Z_]+)", re.IGNORECASE)

#: `@max: 20` · un tope por si un proyecto se dispara y nadie quiere un informe
#: de doscientas diapositivas.
TOPE = re.compile(r"@max\s*:\s*(\d+)", re.IGNORECASE)


def notas_de(slide: Slide) -> str:
    """El texto de las notas del orador, o cadena vacía.

    Una diapositiva puede no tener notas, y preguntar por ellas las crea en
    algunas versiones de `python-pptx`: se comprueba antes de tocar nada para
    no modificar la plantilla solo por leerla.
    """
    if not slide.has_notes_slide:
        return ""
    marco = slide.notes_slide.notes_text_frame
    return marco.text if marco is not None else ""


def plan_de(slide: Slide) -> tuple[str, int] | None:
    """`(colección, tope)` si la diapositiva se repite, o `None`.

    Un `@repeat` sobre una colección que no existe **no se ignora en silencio**:
    devuelve `None` y quien llama lo apunta como aviso. Ignorarlo dejaría una
    diapositiva con los marcadores del primer elemento y aspecto de estar bien.
    """
    notas = notas_de(slide)
    if not notas:
        return None
    encontrado = REPETIR.search(notas)
    if encontrado is None:
        return None
    coleccion = encontrado.group(1).lower()
    if coleccion not in mk.COLECCIONES:
        return None
    tope = TOPE.search(notas)
    return coleccion, int(tope.group(1)) if tope else 0


def mover_detras(prs: Presentacion, slide: Slide, referencia: Slide) -> None:
    """Coloca `slide` justo detrás de `referencia` en el orden de la lista.

    `python-pptx` no ofrece reordenar: se opera sobre `sldIdLst`, que es la
    lista de identificadores del XML de la presentación y es lo que fija el
    orden real. Clonar añade siempre al final, así que sin esto las seis
    diapositivas de activo salen detrás de las conclusiones.
    """
    lista = prs.slides._sldIdLst
    elementos = list(lista)
    por_id = {int(e.get("id")): e for e in elementos}
    elemento = por_id.get(slide.slide_id)
    ancla = por_id.get(referencia.slide_id)
    if elemento is None or ancla is None:
        return
    lista.remove(elemento)
    lista.insert(list(lista).index(ancla) + 1, elemento)


#: El atributo que lleva el identificador de relación de una diapositiva.
_RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def retirar(prs: Presentacion, slide: Slide) -> None:
    """Saca la diapositiva modelo de la presentación, **y renumera el resto**.

    Renumerar no es cosmético: `python-pptx` bautiza cada diapositiva nueva como
    `slide<N>.xml` con **N = cuántas hay**, no con el primer número libre. Quitar
    una de tres y añadir otra produce dos partes llamadas `slide3.xml`, el ZIP
    guarda las dos con el mismo nombre y al abrirlo **una se ha comido a la
    otra**: la presentación sale con la última repetida y la anterior
    desaparecida, sin ningún error por el camino.

    Se vio en el primer informe generado de verdad: dos activos, salía uno.
    `rename_slide_parts` es de la propia biblioteca y deja la numeración
    contigua, que es lo que su contador da por supuesto.
    """
    lista = prs.slides._sldIdLst
    for elemento in list(lista):
        if int(elemento.get("id")) == slide.slide_id:
            prs.part.drop_rel(elemento.get(_RID))
            lista.remove(elemento)
            break
    else:
        return
    prs.part.rename_slide_parts([e.get(_RID) for e in prs.slides._sldIdLst])


def expandir(
    prs: Presentacion,
    snapshot: dict[str, Any],
    globales: dict[str, str],
    *,
    sustituir: Any,
) -> tuple[list[str], list[str]]:
    """Clona las diapositivas modelo y rellena cada copia con su elemento.

    `sustituir(slide, valores)` lo pasa el generador para no duplicar aquí la
    medición de desbordamiento. Devuelve `(marcadores_sin_resolver, avisos)`.

    Las diapositivas fijas **no se tocan**: las rellena el generador con los
    valores globales, como antes. Aquí solo se atienden las que piden repetirse.
    """
    sin_resolver: list[str] = []
    avisos: list[str] = []

    # La lista se congela antes de empezar: clonar añade diapositivas, y
    # recorrer mientras se añade es un bucle que no termina.
    modelos = [(slide, plan_de(slide)) for slide in list(prs.slides)]
    for slide, plan in modelos:
        notas = notas_de(slide)
        if plan is None:
            if REPETIR.search(notas):
                pedida = REPETIR.search(notas)
                avisos.append(
                    f"Una diapositiva pide repetirse sobre «{pedida.group(1) if pedida else '?'}», "
                    f"que no es una colección del informe. Disponibles: "
                    f"{', '.join(sorted(mk.COLECCIONES))}. Se ha dejado sin repetir."
                )
            continue

        coleccion, tope = plan
        clave, valores_de = mk.COLECCIONES[coleccion]
        elementos = list(snapshot.get(clave, []))
        if tope and len(elementos) > tope:
            avisos.append(
                f"«{coleccion}» tiene {len(elementos)} elementos y la plantilla pone un tope de "
                f"{tope}: se han dejado fuera {len(elementos) - tope}."
            )
            elementos = elementos[:tope]

        if not elementos:
            avisos.append(
                f"No hay ningún elemento de «{coleccion}» en este informe, así que su diapositiva "
                "se ha retirado en vez de salir con los marcadores a la vista."
            )
            retirar(prs, slide)
            continue

        anterior: Slide = slide
        for elemento in elementos:
            copia = clonar_diapositiva(prs, slide)
            mover_detras(prs, copia, anterior)
            anterior = copia
            sin_resolver += sustituir(copia, {**globales, **valores_de(snapshot, elemento)})
        retirar(prs, slide)

    return sin_resolver, avisos
