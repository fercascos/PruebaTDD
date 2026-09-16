"""Lo que se coloca **dentro** de la plantilla y no es texto `[REQ]` §3.2.

Dos cosas que el cliente pidió sobre su Full Report con estas palabras:

> *«Todas las fotos que vayamos adjuntando en la parte de inventario deberán
> aparecer en cada recuadro azul, y el pie de cada foto —lo que aparece en la
> plantilla como "Descripción"— es el título de esa foto.»*
>
> *«Sobre el CAPEX, en vez de la tabla que aparece ahí deberá ir la tabla pegada
> de CAPEX de nuestra herramienta.»*

Las dos se declaran en las **notas del orador**, que no se imprimen:

| En las notas | Qué hace |
|---|---|
| `@fotos: HC.H02` | Reparte las fotos de esa sección entre los marcos de la diapositiva |
| `@capex` | Pone la tabla nativa **aquí**, en lugar de lo que hubiera |

## Por qué no hay marcadores para esto

Un marcador es texto que se sustituye por texto. Una fotografía ocupa un marco
con una proporción, y la tabla del CAPEX son diez columnas con subtotales: ni
una ni otra caben en un `{{...}}`. Lo que la plantilla aporta es **la geometría**
—dónde va cada cosa y de qué tamaño—, y eso ya está dibujado en sus
diapositivas: se lee de ahí en vez de pedir que alguien lo escriba.

## Lo que se respeta de la plantilla

**La proporción de la foto.** El marco es 3,15 × 2,36 in —cuatro tercios— y una
fotografía de móvil no lo es: se encaja dentro sin deformarla y se centra. Una
fachada estirada desacredita el informe antes de que nadie lea una línea.

`[LIM]` **Los marcos que sobran se retiran, con su pie.** Una sección con dos
fotos deja dos rectángulos azules vacíos si no se quitan, y un rectángulo azul
vacío en un entregable se lee como un fallo de maquetación.
"""

from __future__ import annotations

import io
import re
from typing import Any

from pptx.presentation import Presentation as Presentacion
from pptx.slide import Slide
from pptx.util import Emu, Inches

from tdd.reporting.marcadores import codigo_y_ancestros
from tdd.reporting.repeticion import mover_detras, notas_de, retirar

#: `@fotos: HC.H02` en las notas de la diapositiva de fotografías.
#:
#: Admite **varios códigos separados por coma**, porque la plantilla del cliente
#: tiene páginas de fotos que sirven a dos secciones: la de ascensores y
#: fontanería, que comparten diapositiva de texto, comparten también la de
#: fotos. Con un solo código, las fotos de la segunda sección se perdían.
FOTOS = re.compile(r"@fotos\s*:\s*([A-Za-z0-9.,\s]+)", re.IGNORECASE)

#: `@capex` en las notas de la diapositiva que debe llevar la tabla, con la
#: lista de qué tablas, separadas por coma. `@capex` a secas es `@capex: detalle`.
#:
#: `[REQ]` La sección 07 de la plantilla del cliente tiene **cinco** tablas
#: distintas, no una: las dos de detalle —obra de arquitectura y de
#: instalaciones—, la matriz de riesgo por plazo que va detrás de cada una, el
#: resumen por capítulo y el presupuesto de costes duros y blandos.
CAPEX = re.compile(r"@capex\b(?:\s*:\s*([A-Za-z:,\s]+))?", re.IGNORECASE)

#: Las tablas que se pueden pedir. `detalle` acepta un bloque detrás de dos
#: puntos —`detalle:arquitectura`— y `riesgos` también, para la matriz que va
#: debajo de cada tabla de detalle.
TABLAS = ("detalle", "riesgos", "capitulos", "costes")

#: Un marco de fotografía de la plantilla del cliente mide 3,15 × 2,36 in. Se
#: admite un margen amplio porque las cuatro plantillas difieren en centésimas.
MARCO_MIN_IN = 1.5

#: Un pie de foto mide 4,11 × 0,27 in: ancho y muy bajo. Es lo que lo distingue
#: del marco, que es alto. No se busca por el texto «Descripción» porque las
#: plantillas inglesas lo escriben en inglés en algunas diapositivas.
PIE_ALTO_MAX_IN = 0.6


def codigos_de_fotos(slide: Slide) -> list[str]:
    """Los códigos que esta diapositiva de fotos reclama. Vacío si no es una."""
    encontrado = FOTOS.search(notas_de(slide))
    if encontrado is None:
        return []
    return [c.strip().upper() for c in encontrado.group(1).split(",") if c.strip()]


def lleva_capex(slide: Slide) -> bool:
    return CAPEX.search(notas_de(slide)) is not None


def tablas_de_capex(slide: Slide) -> list[str]:
    """Qué tablas pide esta diapositiva, en el orden en que se escribieron.

    Vacío si no lleva la directiva. `@capex` sin lista es `["detalle"]`, que es
    lo que significaba antes de que la directiva admitiese nombres: una
    plantilla ya marcada sigue generando lo mismo.
    """
    encontrado = CAPEX.search(notas_de(slide))
    if encontrado is None:
        return []
    lista = encontrado.group(1)
    if not lista:
        return ["detalle"]
    return [t.strip().lower() for t in lista.split(",") if t.strip()]


def secciones_declaradas(prs: Presentacion) -> set[str]:
    """Los códigos que la plantilla reclama con `@fotos`."""
    return {c for s in prs.slides for c in codigos_de_fotos(s)}


def cae_en_una_seccion(foto: Any, declaradas: set[str]) -> bool:
    """¿Va esta fotografía a alguna diapositiva de sección?

    Se mira contra el código de la foto **y sus ancestros**: la de una
    enfriadora `HC.H08.01` cae en la diapositiva de climatización, que pide el
    capítulo `HC.H08`.
    """
    codigo = getattr(foto, "capex_code", "") or ""
    return bool(declaradas & set(codigo_y_ancestros(codigo))) if codigo else False


def _marcos_y_pies(slide: Slide) -> tuple[list[Any], list[Any]]:
    """Los huecos de imagen y sus pies, **en el orden de lectura**.

    De arriba abajo y de izquierda a derecha, que es como se miran cuatro fotos
    en una diapositiva. Sin ordenar, el orden es el del XML —el de creación en
    PowerPoint— y la tercera foto podía salir arriba a la derecha.
    """
    marcos, pies = [], []
    for forma in slide.shapes:
        if forma.width is None or forma.height is None:
            continue
        alto = Emu(forma.height).inches
        ancho = Emu(forma.width).inches
        if alto >= MARCO_MIN_IN and ancho >= MARCO_MIN_IN:
            marcos.append(forma)
        elif alto <= PIE_ALTO_MAX_IN and ancho >= MARCO_MIN_IN and forma.has_text_frame:
            pies.append(forma)

    def por_lectura(f: Any) -> tuple[int, int]:
        # Redondeado a un cuarto de pulgada: dos marcos de la misma fila no
        # están exactamente a la misma altura en la plantilla real —difieren en
        # centésimas— y sin redondear se ordenarían en escalera.
        return (round(Emu(f.top).inches * 4), round(Emu(f.left).inches * 4))

    return sorted(marcos, key=por_lectura), sorted(pies, key=por_lectura)


def _fotos_de(fotos: list[Any], codigos: list[str]) -> list[Any]:
    """Las fotografías de una o varias secciones, en el orden del consultor.

    Una foto entra en la sección de su propio código y en la de sus ancestros:
    la de una enfriadora —`HC.H08.01`— sale en la diapositiva de climatización
    aunque esa esté atada al capítulo `HC.H08`.
    """
    quiere = set(codigos)
    return [
        f for f in fotos if quiere & set(codigo_y_ancestros(getattr(f, "capex_code", "") or ""))
    ]


def _colocar(slide: Slide, marco: Any, datos: bytes) -> bool:
    """Pone la imagen en el hueco del marco y retira el marco.

    Se conserva la proporción: se encaja dentro y se centra. El marco original
    se retira después, porque es un rectángulo de color que se vería detrás.
    """
    from PIL import Image

    try:
        with Image.open(io.BytesIO(datos)) as img:
            ancho_px, alto_px = img.size
    except Exception:  # noqa: BLE001 — una foto ilegible no tumba el informe
        return False
    proporcion = ancho_px / alto_px if alto_px else 1.0

    hueco_ancho = Emu(marco.width).inches
    hueco_alto = Emu(marco.height).inches
    ancho = hueco_ancho
    alto = ancho / proporcion
    if alto > hueco_alto:
        alto = hueco_alto
        ancho = alto * proporcion

    izquierda = Emu(marco.left).inches + (hueco_ancho - ancho) / 2
    arriba = Emu(marco.top).inches + (hueco_alto - alto) / 2
    slide.shapes.add_picture(
        io.BytesIO(datos), Inches(izquierda), Inches(arriba), Inches(ancho), Inches(alto)
    )
    marco._element.getparent().remove(marco._element)
    return True


def _titulo(foto: Any) -> str:
    """El pie: el título que le puso quien la subió.

    `[REQ]` Con las palabras del cliente: *«el pie de cada foto, lo que aparece
    en la plantilla como "Descripción", es el título de esa foto»*.
    """
    return str(getattr(foto, "caption", "") or "")


def _vaciar(forma: Any) -> None:
    forma._element.getparent().remove(forma._element)


def repartir_fotos(prs: Presentacion, fotos: list[Any], *, clonar: Any) -> tuple[int, list[str]]:
    """Coloca las fotos de cada sección en su diapositiva. Devuelve `(puestas, avisos)`.

    Con más fotos que marcos, la diapositiva **se clona**: cuatro por página,
    tantas páginas como hagan falta. Es lo que hacía a mano quien montaba el
    informe, y la plantilla ya trae la maqueta de esa página.
    """
    puestas = 0
    avisos: list[str] = []
    sin_fotos: list[str] = []
    for slide in list(prs.slides):
        codigos = codigos_de_fotos(slide)
        if not codigos:
            continue
        codigo = " + ".join(codigos)
        marcos, pies = _marcos_y_pies(slide)
        if not marcos:
            avisos.append(
                f"La diapositiva de fotos de «{codigo}» no tiene ningún marco reconocible: "
                "se esperan formas de al menos 1,5 × 1,5 in."
            )
            continue

        suyas = _fotos_de(fotos, codigos)
        if not suyas:
            # `[REQ]` Sin fotos **se retira la diapositiva entera**, no solo sus
            # marcos. Vaciarlos dejaba una página con la cabecera, el pie y nada
            # más en medio, y de esas salían diez en un informe de sesenta y
            # nueve: en un entregable eso parece un fallo de impresión. Es la
            # misma regla que ya seguía `@repeat` con una colección vacía.
            #
            # `[LIM]` Si lo que se quiere es el hueco para pegar fotos a mano
            # después, se le quita el `@fotos` a esa diapositiva en las notas.
            retirar(prs, slide)
            sin_fotos.append(codigo)
            continue

        paginas = [suyas[i : i + len(marcos)] for i in range(0, len(suyas), len(marcos))]
        destino: Slide = slide
        for indice, tanda in enumerate(paginas):
            if indice > 0:
                destino = clonar(prs, slide)
            _rellenar(destino, tanda)
            puestas += len(tanda)
        if len(paginas) > 1:
            avisos.append(
                f"«{codigo}» tiene {len(suyas)} fotografías y caben {len(marcos)} por "
                f"diapositiva: se han añadido {len(paginas) - 1} más."
            )

    if sin_fotos:
        avisos.append(
            f"{len(sin_fotos)} secciones no tienen ninguna fotografía y su diapositiva se ha "
            f"retirado en vez de salir en blanco: {', '.join(sin_fotos)}."
        )
    return puestas, avisos


def _rellenar(slide: Slide, tanda: list[Any]) -> None:
    """Una tanda de fotos en los marcos de UNA diapositiva."""
    marcos, pies = _marcos_y_pies(slide)
    for indice, marco in enumerate(marcos):
        pie = pies[indice] if indice < len(pies) else None
        if indice < len(tanda):
            if _colocar(slide, marco, tanda[indice].datos) and pie is not None:
                _escribir_pie(pie, _titulo(tanda[indice]))
        else:
            _vaciar(marco)
            if pie is not None:
                _vaciar(pie)


def _escribir_pie(forma: Any, texto: str) -> None:
    """El título de la foto, **con el formato del pie de la plantilla**.

    Se escribe en el primer `run` y se borran los demás, que es lo mismo que
    hace la sustitución de marcadores: vaciar el marco y añadir texto suelto le
    quitaría la tipografía y la cursiva que la plantilla le puso.
    """
    parrafos = forma.text_frame.paragraphs
    if not parrafos or not parrafos[0].runs:
        forma.text_frame.text = texto
        return
    parrafos[0].runs[0].text = texto
    for sobrante in parrafos[0].runs[1:]:
        sobrante._r.getparent().remove(sobrante._r)
    for otro in parrafos[1:]:
        otro._p.getparent().remove(otro._p)


#: Dónde arranca la tabla de detalle cuando ocupa la diapositiva entera. Es la
#: posición que tiene en la plantilla del cliente, medida sobre su fichero.
DETALLE_IZQ_IN = 0.47
DETALLE_ARRIBA_IN = 1.55


def _imagenes(slide: Slide) -> list[Any]:
    """Las imágenes de la diapositiva, **de mayor a menor superficie**.

    La más grande es la tabla pegada desde Excel; las pequeñas de las
    diapositivas 56 y 58 son la **leyenda de riesgo**, que es contenido de la
    plantilla y no se toca. Retirarlas todas, como se hacía, dejaba la página de
    la matriz sin sus cuatro cuadros de color ni la escala de plazos.
    """
    fotos = [
        f
        for f in slide.shapes
        if f.shape_type == 13 and f.width is not None and f.height is not None  # noqa: PLR2004
    ]
    return sorted(fotos, key=lambda f: -(Emu(f.width).inches * Emu(f.height).inches))


def poner_capex(
    prs: Presentacion, tablas: dict[str, list[Any]], *, clonar: Any, insertar: Any
) -> tuple[int, list[str]]:
    """Pone las tablas nativas donde la plantilla las pide. `(puestas, avisos)`.

    `[REQ]` Con las palabras del cliente: *«en vez de la tabla que aparece ahí
    deberá ir la tabla pegada de CAPEX de nuestra herramienta»*. Lo que había en
    esas diapositivas son **imágenes** pegadas desde Excel, con los números de
    otro proyecto: se retiran, porque dejarlas debajo daría dos tablas con
    cifras distintas en la misma página.

    `tablas` es `{nombre: [trozos]}`. Cada diapositiva declara en sus notas qué
    tablas quiere, y cada una ocupa el marco de la imagen que sustituye:

        @capex: detalle:arquitectura     la tabla de obra de arquitectura
        @capex: riesgos:arquitectura     su matriz de riesgo por plazo
        @capex: capitulos, riesgos       las dos, apiladas en el mismo hueco
        @capex: costes                   el presupuesto de duros y blandos

    Una tabla que no cabe se parte, y cada trozo se lleva su copia de la
    diapositiva —con su cabecera y su pie— detrás de la anterior.
    """
    avisos: list[str] = []
    puestas = 0

    for slide in list(prs.slides):
        pedidas = tablas_de_capex(slide)
        if not pedidas:
            continue
        desconocidas = [p for p in pedidas if p.split(":")[0] not in TABLAS]
        if desconocidas:
            avisos.append(
                f"«@capex: {', '.join(desconocidas)}» no es una tabla del informe. "
                f"Disponibles: {', '.join(TABLAS)}. Esa diapositiva se ha dejado como estaba."
            )
            continue
        conocidas = [p for p in pedidas if p in tablas]
        faltan = [p for p in pedidas if p not in tablas]
        if faltan:
            avisos.append(
                f"La plantilla pide «{', '.join(faltan)}» y este proyecto no tiene datos para "
                "esa tabla. Se ha dejado el hueco en vez de una tabla vacía."
            )
        if not conocidas:
            continue

        puestas += _poner_en(
            prs, slide, [(n, tablas[n]) for n in conocidas], clonar=clonar, insertar=insertar
        )

    return puestas, avisos


def _poner_en(
    prs: Presentacion,
    modelo: Slide,
    pedidas: list[tuple[str, list[Any]]],
    *,
    clonar: Any,
    insertar: Any,
) -> int:
    """Dibuja en una diapositiva las tablas que ha pedido. Devuelve cuántas."""
    marcos = _imagenes(modelo)
    # Una sola tabla que ocupa la página entera —las dos de detalle— se dibuja
    # en la posición de la plantilla y se lleva por delante todas sus imágenes,
    # que son los trozos de la tabla vieja. Con varias tablas, cada una va al
    # marco de la imagen que sustituye y las demás se quedan.
    a_pagina_completa = len(pedidas) == 1 and pedidas[0][0].startswith("detalle")

    # Las copias se sacan **antes** de escribir nada, y todas del modelo intacto.
    # Clonarlo sobre la marcha lo copiaba con el trozo anterior ya dentro: el
    # «(2/2)» salía dibujado encima de las filas del «(1/2)».
    trozos_extra = max((len(t) for _, t in pedidas), default=1) - 1
    destinos: list[Slide] = [modelo]
    anterior: Slide = modelo
    for _ in range(trozos_extra):
        copia = clonar(prs, modelo)
        # Clonar añade al final de la presentación: sin esto el «(2/2)» salía
        # como última diapositiva del informe, detrás de las conclusiones.
        mover_detras(prs, copia, anterior)
        destinos.append(copia)
        anterior = copia

    puestas = 0
    for indice, destino in enumerate(destinos):
        if a_pagina_completa:
            for forma in _imagenes(destino):
                _vaciar(forma)

        arriba = None
        for orden, (_, trozos) in enumerate(pedidas):
            if indice >= len(trozos):
                continue
            if a_pagina_completa:
                izq, arr, ancho = DETALLE_IZQ_IN, DETALLE_ARRIBA_IN, None
            else:
                # El marco de esta tabla es la imagen que le toca por tamaño; si
                # dos tablas comparten un solo marco —el resumen por capítulo y
                # la matriz de la diapositiva 59 vienen en una sola imagen— la
                # segunda se apila debajo de la primera.
                marco = marcos[min(orden, len(marcos) - 1)] if marcos else None
                if marco is None:
                    continue
                izq = Emu(marco.left).inches
                arr = arriba if orden and arriba is not None else Emu(marco.top).inches
                ancho = Emu(marco.width).inches
                if indice == 0 and orden < len(marcos):
                    _vaciar(marco)
            alto = insertar(destino, trozos[indice], izq, arr, ancho)
            arriba = (arr + alto) if alto else arr
            puestas += 1
    return puestas
