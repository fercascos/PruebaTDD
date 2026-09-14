"""Una plantilla de informe de ejemplo, con marcadores y `@repeat`.

    python3 tools/plantilla_de_ejemplo.py ejemplo.pptx

`[REQ]` Sirve para **dos cosas**, y las dos hacen falta:

1. **Enseñar la sintaxis sin leer documentación.** Se abre en PowerPoint y se ve
   dónde va cada `{{marcador}}` y cómo se escribe `@repeat` en las notas del
   orador. Copiar de un fichero que funciona es más rápido que seguir una tabla.
2. **Comprobar que la instalación genera.** Si esta plantilla no produce un
   informe con datos dentro, el problema no está en la plantilla de nadie.

`[LIM]` **No es un diseño.** Es texto sobre fondo blanco a propósito: la
plantilla de verdad la pone el cliente con su marca, y un ejemplo con aspecto de
acabado invita a usarlo tal cual. Aquí lo único que importa es dónde van los
marcadores.

`[REQ]` §15 · Ningún dato real: los marcadores se sustituyen al generar.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

#: `(título, cuerpo, notas)`. El cuerpo lleva los marcadores.
DIAPOSITIVAS: tuple[tuple[str, str, str], ...] = (
    (
        "{{project.name}}",
        "Referencia: {{project.code}}\n"
        "Cliente: {{client.name}}\n"
        "Fecha del informe: {{report.date}}\n"
        "{{project.asset_count}} activos · {{project.finding_count}} hallazgos",
        "",
    ),
    (
        "Resumen económico",
        "CAPEX total: {{capex.total}}\n\n"
        "Por plazo\n"
        "  Corto: {{capex.corto}}\n"
        "  Medio: {{capex.medio}}\n"
        "  Largo: {{capex.largo}}\n"
        "  Mejoras: {{capex.mejoras}}\n\n"
        "Quién lo paga\n"
        "  Lo asume la propiedad: {{capex.propiedad}}\n"
        "  Repercutible al inquilino: {{capex.inquilino}}\n"
        "  Sin determinar: {{capex.sin_determinar}}",
        "",
    ),
    (
        "{{asset.name}}",
        "{{asset.address}}\n"
        "Tipología: {{asset.typology}} · Construido en {{asset.year_built}}\n"
        "Superficie construida: {{asset.total_built_sqm}} m²\n"
        "Visita: {{asset.visit_date}}\n\n"
        "CAPEX del activo: {{asset.capex_total}} "
        "(propiedad {{asset.capex_propiedad}} · inquilino {{asset.capex_inquilino}})\n\n"
        "Descriptivo\n{{asset.descriptivo}}\n\n"
        "Valoración\n{{asset.valoracion}}",
        # Una diapositiva por activo. La instrucción va en las notas porque NO
        # se imprime: la diapositiva se ve igual en la plantilla y en el informe.
        "@repeat: asset",
    ),
    (
        "{{finding.risk_code}} · {{finding.title}}",
        "Zona: {{finding.zone}}\n"
        "Capítulo: {{finding.capex_chapter}} · {{finding.capex_item}}\n"
        "Concepto: {{finding.concept}}\n\n"
        "{{finding.description}}\n\n"
        "Recomendación: {{finding.recommendation}}\n\n"
        "Importe: {{finding.amount}} ({{finding.horizons}}) · Paga: {{finding.payer}}\n"
        "Riesgo {{finding.risk_name}}: {{finding.risk_definition}}",
        # El tope evita un informe de doscientas diapositivas si el proyecto se
        # dispara. Lo que quede fuera se avisa al generar, no se pierde callando.
        "@repeat: finding\n@max: 25",
    ),
    (
        "Limitaciones y salvedades",
        "{{report.limitations}}\n\n"
        "Limitaciones de acceso durante las visitas:\n{{visit.access_limitations}}",
        "",
    ),
    (
        "Criterio de riesgo",
        "{{risk.legend}}\n\n"
        "Extremo: {{risk.04}} · Alto: {{risk.03}} · "
        "Moderado: {{risk.02}} · Bajo: {{risk.01}}",
        "",
    ),
)


def construir() -> Presentation:
    prs = Presentation()
    # 16:9, que es lo que usa cualquier plantilla corporativa de hoy.
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for titulo, cuerpo, notas in DIAPOSITIVAS:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        caja = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
        caja.text_frame.text = titulo
        caja.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        caja.text_frame.paragraphs[0].runs[0].font.bold = True

        texto = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
        texto.text_frame.word_wrap = True
        texto.text_frame.text = cuerpo
        for parrafo in texto.text_frame.paragraphs:
            for run in parrafo.runs:
                run.font.size = Pt(14)

        if notas:
            slide.notes_slide.notes_text_frame.text = notas
    return prs


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("salida", type=Path, help="el .pptx que se escribe")
    args = parser.parse_args(argv)

    # No se sobrescribe sin querer: una plantilla que alguien haya editado a
    # mano encima de esta se perdería sin aviso.
    if args.salida.exists():
        print(f"{args.salida} ya existe. Elija otro nombre o bórrelo.", file=sys.stderr)
        return 2

    construir().save(args.salida)
    repetidas = sum(1 for _, _, notas in DIAPOSITIVAS if "@repeat" in notas)
    print(
        f"{args.salida} · {len(DIAPOSITIVAS)} diapositivas, {repetidas} de ellas repetibles.\n"
        "Ábrala en PowerPoint: los marcadores están en el cuerpo y las directivas "
        "en las notas del orador."
    )
    return 0


if __name__ == "__main__":  # pragma: no cover — punto de entrada
    raise SystemExit(main())
